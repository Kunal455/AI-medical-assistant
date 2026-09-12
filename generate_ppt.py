from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title_text, bullet_points):
    slide_layout = prs.slide_layouts[1] # title and content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    for i, bp in enumerate(bullet_points):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bp
        p.level = 0
        p.font.size = Pt(24)

# Slide 1: Title
slide_layout = prs.slide_layouts[0] # title slide
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "DevOps Implementation for AI Medical Assistant"
slide.shapes.placeholders[1].text = "Containerization, CI/CD, and Microservices Orchestration"

# Slide 2
add_slide("Project Overview (A DevOps Perspective)", [
    "A full-stack web application delivering AI-driven medical symptom triage.",
    "The DevOps Challenge: Managing a complex, polyglot architecture.",
    "Three distinct microservices: React.js, Node.js, Python/FastAPI.",
    "The DevOps Goal: Ensure consistent environments, automate CI/CD, zero-downtime."
])

# Slide 3
add_slide("The DevOps Tech Stack", [
    "Version Control: Git & GitHub (Branching, Webhooks).",
    "Containerization: Docker (Isolated application environments).",
    "Orchestration: Docker Compose (Multi-container networking).",
    "CI/CD Automation: Jenkins (Declarative Pipelines).",
    "Infrastructure: Linux/Ubuntu Server (Host environment)."
])

# Slide 4
add_slide("Microservices Architecture", [
    "Shifted from Monolithic architecture to Microservices.",
    "Separation of Concerns:",
    "  - UI updates don't require AI backend reboots.",
    "  - API Gateway (Node) handles load balancing and security.",
    "  - Heavy computation (Python) is isolated.",
    "Benefits: Independent scaling, targeted debugging, polyglot programming."
])

# Slide 5
add_slide("Containerization Strategy (Docker)", [
    "Created distinct Dockerfiles for each service.",
    "Frontend Dockerfile: Node Alpine image, exposes port 3000.",
    "Node Backend Dockerfile: Node environment, Express routing on port 5000.",
    "Python Backend Dockerfile: Python 3.9 and ML libraries, port 8000.",
    "Result: 'It works on my machine' syndrome is completely eliminated."
])

# Slide 6
add_slide("Microservice Orchestration (Docker Compose)", [
    "docker-compose.yml serves as the blueprint for the infrastructure.",
    "Custom Bridge Network: medassist_network isolates traffic.",
    "Internal DNS Resolution: Containers communicate via names (e.g., python-backend).",
    "Volume Mounting: Local directories mapped for hot-reloading in dev."
])

# Slide 7
add_slide("Continuous Integration / Deployment (Jenkins)", [
    "Implemented a fully automated CI/CD pipeline.",
    "Trigger: GitHub Webhooks automatically notify Jenkins upon 'push'.",
    "Pipeline-as-Code: Jenkinsfile stored directly in the repository.",
    "Goal: Catch bugs early, build instantly, deploy seamlessly."
])

# Slide 8
add_slide("The Jenkins Pipeline Stages", [
    "1. Checkout Stage: Pulls latest commit from GitHub.",
    "2. Linting & Security: npm run lint and flake8.",
    "3. Dependency Installation: npm install and pip install.",
    "4. Build Stage: docker-compose build --no-cache.",
    "5. Deploy Stage: docker-compose up -d to restart the application."
])

# Slide 9
add_slide("Results & Impact of DevOps Integration", [
    "Faster Release Cycles: Automatic updates in minutes.",
    "High Reliability: Automated dependency installation prevents crashes.",
    "Resource Efficiency: Alpine base images reduced sizes by over 60%.",
    "Easy Onboarding: docker-compose up spins up everything instantly."
])

# Slide 10
add_slide("Future DevOps Enhancements", [
    "Kubernetes (K8s) Migration: EKS for pod auto-scaling and self-healing.",
    "Observability Stack: Prometheus and Grafana for monitoring.",
    "Log Aggregation: ELK Stack (Elasticsearch, Logstash, Kibana).",
    "Infrastructure as Code (IaC): Terraform for provisioning."
])

# Slide 11
add_slide("Q&A", [
    "Thank you for your time!",
    "Opening the floor for questions regarding Docker, Jenkins, or microservices architecture."
])

prs.save('DevOps_Presentation.pptx')
print("DevOps_Presentation.pptx generated successfully!")
