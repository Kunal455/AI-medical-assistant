from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title_text, bullet_points):
    slide_layout = prs.slide_layouts[1] # title and content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    for i, bp in enumerate(bullet_points):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bp
        p.level = 0
        p.font.size = Pt(22)

# Slide 1: Title
slide_layout = prs.slide_layouts[0] # title slide
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "AI Medical Assistant"
slide.shapes.placeholders[1].text = "An Intelligent Healthcare Triage Platform\n[Your Name]"

# Slide 2: Introduction & Problem Statement
add_slide("Introduction & Problem Statement", [
    "Problem: Long wait times for basic medical consultations.",
    "Problem: Online self-diagnosis often leads to 'cyberchondria'.",
    "Solution: A smart, AI-driven assistant that understands natural language.",
    "Goal: Provide immediate, safe, and preliminary symptom analysis."
])

# Slide 3: Objectives
add_slide("Project Objectives", [
    "Develop an intelligent chatbot for medical queries.",
    "Implement a robust Microservices Architecture.",
    "Ensure secure, low-latency communication between services.",
    "Automate deployment using a robust CI/CD pipeline."
])

# Slide 4: System Architecture & Tech Stack
add_slide("Architecture & Tech Stack", [
    "Frontend: React.js (Modern, responsive UI).",
    "API Gateway: Node.js & Express (Routing, Security).",
    "AI Engine: Python & FastAPI (Machine Learning Inference).",
    "Decoupled microservices ensure independent scaling and fault tolerance."
])

# Slide 5: Workflow - How It Works
add_slide("Workflow", [
    "1. User enters symptoms via the React web interface.",
    "2. Node.js Gateway validates and securely routes the request.",
    "3. Python Backend processes text using an NLP model.",
    "4. AI cross-references data and generates medical advice.",
    "5. Response is instantly relayed back to the user's screen."
])

# Slide 6: DevOps & CI/CD
add_slide("DevOps Implementation", [
    "Containerization: Docker ensures consistent environments.",
    "Orchestration: Docker Compose handles internal networking.",
    "Continuous Integration: Jenkins automates building & testing.",
    "Pipeline: GitHub push triggers instant Jenkins builds.",
    "Result: Zero 'works on my machine' errors."
])

# Slide 7: Results & Interface
add_slide("Results & Interface", [
    "Successfully deployed multi-tier application.",
    "Real-time processing of complex medical symptoms.",
    "[Insert Screenshot of UI here]",
    "[Insert Screenshot of Jenkins Success here]"
])

# Slide 8: Future Scope
add_slide("Future Enhancements", [
    "EHR Integration: Connect to electronic health records (HL7/FHIR).",
    "Multi-Modal AI: Analyze uploaded images (e.g., skin rashes).",
    "Voice-to-Text: Improved accessibility for all users.",
    "Cloud Native: Migrate to Kubernetes for auto-scaling."
])

# Slide 9: Conclusion
add_slide("Conclusion", [
    "Successfully bridged healthcare and modern AI tech.",
    "Proven reliability through microservices and Docker.",
    "A scalable foundation ready for future healthcare integrations.",
    "Thank you! Any questions?"
])

prs.save('Full_Project_Presentation.pptx')
print("Full_Project_Presentation.pptx generated successfully!")
